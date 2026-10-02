from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; SQL=ROOT/'sql'; OUT=ROOT/'docs/database-requirements-ja.docx'
def font(r,sz=8,b=False,c=None):
    r.font.name='Yu Gothic'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Yu Gothic'); r.font.size=Pt(sz); r.bold=b
    if c:r.font.color.rgb=RGBColor(*c)
def text(p,s,sz=8,b=False,c=None): font(p.add_run(str(s)),sz,b,c)
def heading(d,s,l=1):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(8 if l==1 else 4); text(p,s,15 if l==1 else 11,True,(177,55,43) if l==1 else (35,35,35)); return p
def para(d,s,sz=8): p=d.add_paragraph(); text(p,s,sz); return p
def table(d,heads,rows,fs=7):
    t=d.add_table(rows=1,cols=len(heads)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(heads):
        c=t.rows[0].cells[i]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; text(c.paragraphs[0],h,fs,True)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; text(cells[i].paragraphs[0],v,fs)
    d.add_paragraph(); return t

base=(SQL/'01_create_tables.sql').read_text(encoding='utf-8')
tables=re.findall(r'(?im)^\s*CREATE TABLE\s+([A-Z0-9_]+)',base); seqs=re.findall(r'(?im)^\s*CREATE SEQUENCE\s+([A-Z0-9_]+)',base)
def cols(n):
    m=re.search(r'CREATE TABLE\s+'+n+r'\s*\((.*?)(?=\n\);)',base,re.I|re.S); out=[]
    if m:
        for x in m.group(1).splitlines():
            z=re.match(r'\s*([A-Z][A-Z0-9_]*)\s+(.+)',x.rstrip(',').strip(),re.I)
            if z and not z.group(1).upper() in ('CONSTRAINT','PRIMARY','UNIQUE','CHECK','FOREIGN'): out.append(z.groups())
    return out
files=sorted(p.name for p in SQL.glob('*.sql'))
purpose={
'00_check_environment.sql':'既存JL_オブジェクトとOracle環境の確認','01_create_tables.sql':'基本テーブル・制約・インデックス・シーケンス作成','02_sample_data.sql':'東京・港区を起点とする基本案内データ投入','03_create_translation_tables.sql':'多言語テーブルの個別作成','03_verify.sql':'作成件数と主要データの検証','04_edit_examples.sql':'既存案内を編集する参考SQL','04_seed_translation_data.sql':'韓国語・英語翻訳データ投入','05_create_task_events.sql':'利用者タスクイベント構造の補完','06_create_document_guidance.sql':'書類説明・準備方法の補完','07_update_document_guidance.sql':'書類案内の更新','08_add_resident_procedures.sql':'生活行政手続6件の追加','09_review_resident_procedure_documents.sql':'追加手続の書類・翻訳・条件補完','10_add_email_verification.sql':'メール認証構造の追加','11_add_password_reset.sql':'パスワード再設定トークンの追加','12_add_member_status.sql':'会員状態・認証状態の追加','13_create_community.sql':'投稿・コメント・いいね・通報の追加','14_add_community_image.sql':'コミュニティ画像構造の追加','15_seed_community_samples.sql':'コミュニティ確認用サンプル投入','16_add_community_view_count.sql':'投稿閲覧数の追加','16_preview_life_procedures.sql':'生活行政手続の追加・プレビュー投入','17_update_life_procedure_document_guidance.sql':'生活手続の書類案内更新','18_update_life_document_condition_translations.sql':'条件付き書類の翻訳更新','19_update_life_procedure_offices.sql':'手続と窓口の関連更新','20_fix_japanese_life_procedure_steps.sql':'日本語手順本文修正','21_fix_japanese_life_procedure_step_titles.sql':'日本語手順タイトル修正','22_fix_japanese_life_procedure_overview.sql':'日本語手続概要修正','23_fix_japanese_life_document_conditions.sql':'日本語書類条件修正','24_fix_japanese_life_procedure_overview_11_15.sql':'手続11〜15の日本語概要修正','25_create_community_user_features.sql':'ブックマーク・通知・監査・コメント更新日時の補完','install_all.sql':'初回インストール実行順の統合'}
d=Document(); s=d.sections[0]; s.top_margin=Inches(.65); s.bottom_margin=Inches(.65); s.left_margin=Inches(.68); s.right_margin=Inches(.68)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,'Japan Life Guide',18,True,(177,55,43))
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,'データベース要件定義書',23,True)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,'Javaフルスタック 個人プロジェクト成果物',10,False,(100,90,75))
para(d,'対象システム：東京都港区の生活行政案内・手続管理・コミュニティ機能'); para(d,'根拠：プロジェクトのsqlフォルダ内にある全SQLファイル、docs/database.md、README.md。')
heading(d,'1. 文書の目的と範囲'); para(d,'本書は、Japan Life Guideで使用するOracleデータベースのデータ要件、テーブル構成、参照関係、状態値、初期投入・更新SQLを整理したものです。SQLから確認できるデータベース実装範囲を対象とします。')
table(d,['項目','内容'],[['目的','生活行政手続、必要書類、窓口、利用者の進捗、コミュニティ情報を管理する。'],['対象DB','Oracle。通常のプロジェクト用アカウントで実行する前提。'],['文字データ','日本語投入ではN\'…\'形式を使用するSQLがある。'],['対象外','アカウント作成、既存オブジェクト一括削除、SYS/SYSTEM向け操作。']])
heading(d,'2. システムとデータベースの概要'); para(d,'地域・カテゴリを起点に生活行政手続を表示し、手続ごとの書類、手順、窓口、情報源を管理します。会員は手続を保存してチェック状態を管理でき、コミュニティでは投稿、コメント、いいね、画像、ブックマーク、通報を扱います。')
table(d,['領域','主なデータ','テーブル'],[['生活行政案内','手続・書類・手順・情報源・窓口','JL_PROCEDURE、JL_DOCUMENT、JL_STEP、JL_SOURCE、JL_OFFICE'],['多言語','手続・書類・手順・情報源・窓口の翻訳','各 *_I18N、JL_DOCUMENT_GUIDE'],['会員・認証','会員、メール認証、パスワード再設定','JL_MEMBER、JL_EMAIL_VERIFICATION、JL_PASSWORD_RESET'],['利用者進捗','保存手続、チェック項目、イベント','JL_USER_TASK、JL_USER_CHECK、JL_USER_TASK_EVENT'],['コミュニティ','投稿、コメント、画像、反応、通報','JL_COMMUNITY_*'],['運営','通知、会員監査、通報監査','JL_NOTIFICATION、JL_MEMBER_AUDIT、JL_COMMUNITY_REPORT_AUDIT']])
heading(d,'3. 実装範囲'); table(d,['機能','データ要件','根拠SQL'],[['手続案内','公開状態、状況区分、段階、改訂番号を持つ。','01、02、08、16_preview、17〜24'],['書類・手順','表示順、必須/条件付き、説明、準備方法、翻訳を管理する。','01、04、07、09、17、18、23'],['窓口','地域・窓口情報を持ち、手続と関連する。','01、02、08、19'],['認証','メール認証とパスワード再設定トークンを管理する。','10、11、12'],['タスク','会員ごとの保存・チェック・イベントを保持する。','01、05'],['コミュニティ','投稿・コメント・反応・画像・通報・運営処理を管理する。','13、14、16、25']])
heading(d,'4. 機能別データ要件'); table(d,['ID','要件','関連テーブル','実装基準'],[['D-01','地域・カテゴリをマスタとして保持する。','JL_REGION、JL_CATEGORY','手続や窓口から参照される親データ。'],['D-02','手続の公開可否と改訂番号を管理する。','JL_PROCEDURE','PUBLISHED、REVISION等の制約で値を限定。'],['D-03','書類の必須/条件付き状態を管理する。','JL_DOCUMENT','条件付きの場合は条件説明を必要とする。'],['D-04','案内本文を言語別に提供する。','*_I18N、JL_DOCUMENT_GUIDE','language_codeで言語を識別。'],['D-05','会員の保存状態を改訂時点とともに残す。','JL_USER_TASK、JL_USER_CHECK','saved_revisionと完了日時を持つ。'],['D-06','表示状態と通報状態を分離する。','投稿・コメント・通報','ACTIVE/HIDDEN/DELETEDとOPEN/RESOLVED/REJECTEDを用途別に使用。'],['D-07','運営者の通報処理を監査記録に残す。','JL_COMMUNITY_REPORT_AUDIT','前後の通報状態、対象種別、アクションを保存。']])
heading(d,'5. テーブル一覧'); table(d,['No.','テーブル','役割'],[[i,n,'01_create_tables.sqlにCREATE TABLE定義あり'] for i,n in enumerate(tables,1)])
para(d,'現行の01_create_tables.sqlには29テーブルのCREATE TABLE定義があります。多言語テーブルやコミュニティ拡張も含めた数です。')
heading(d,'6. 主要テーブル要件')
for n in ['JL_MEMBER','JL_PROCEDURE','JL_DOCUMENT','JL_STEP','JL_OFFICE','JL_USER_TASK','JL_COMMUNITY_POST','JL_COMMUNITY_COMMENT','JL_COMMUNITY_REPORT','JL_COMMUNITY_REPORT_AUDIT']:
    heading(d,n,2); table(d,['列','定義'],cols(n))
heading(d,'7. 参照関係と削除方針'); para(d,'外部キーは手続・書類・手順・窓口の親子関係、会員の利用者データ、コミュニティの投稿・コメント・通報関係を表します。明示されたON DELETE CASCADEは親削除時の子データ削除に使用され、それ以外は参照整合性を優先します。')
table(d,['親','子','関係'],[['JL_REGION','CATEGORY / PROCEDURE / OFFICE','地域を基準に分類'],['JL_PROCEDURE','DOCUMENT / STEP / SOURCE / PROC_OFFICE','一つの手続に複数の関連データ'],['JL_MEMBER','USER_TASK / POST / NOTIFICATION / AUDIT','会員を基準に進捗・投稿・監査を管理'],['JL_USER_TASK','USER_CHECK / USER_TASK_EVENT','保存手続のチェックとイベント'],['JL_COMMUNITY_POST','COMMENT / LIKE / IMAGE / REPORT / BOOKMARK','投稿中心の関連データ'],['JL_COMMUNITY_REPORT','REPORT_AUDIT','通報状態と処理履歴']])
heading(d,'8. 状態値と管理ログ'); table(d,['項目','許容値','意味'],[['会員ロール','MEMBER / ADMIN','一般会員と管理者'],['会員状態','ACTIVE / INACTIVE','利用可能または無効'],['メール認証','Y / N','認証済みかどうか'],['手続公開','Y / N','表示対象かどうか'],['手続状況','ABROAD / DOMESTIC','生活状況の区分'],['タスク状態','NOT_STARTED / PREPARING / COMPLETED','会員の進捗'],['投稿・コメント','ACTIVE / HIDDEN / DELETED','表示状態'],['通報','OPEN / RESOLVED / REJECTED','未処理・解決・却下'],['通報アクション','REJECT / HIDE / DELETE','運営操作']]); para(d,'通報監査は前後の通報状態、対象種別、アクション、管理者ID、理由、作成日時を保持します。')
heading(d,'9. 多言語データと案内表示'); table(d,['テーブル','用途'],[['JL_PROCEDURE_I18N','手続タイトル・概要等の翻訳'],['JL_DOCUMENT_I18N','書類名・条件説明等の翻訳'],['JL_STEP_I18N','手順タイトル・本文の翻訳'],['JL_SOURCE_I18N','出典表示の翻訳'],['JL_OFFICE_I18N','窓口名称・所在地等の翻訳'],['JL_DOCUMENT_GUIDE','説明・準備方法の言語別案内']]); para(d,'SQLで使用される言語識別値はja・ko・enを対象とします。翻訳がない場合の画面側フォールバックはアプリケーション側の責務です。')
heading(d,'10. シーケンスと採番'); table(d,['No.','シーケンス','用途'],[[i,n,'ID採番'] for i,n in enumerate(seqs,1)]); para(d,'現行の01_create_tables.sqlには20個のOracleシーケンス定義があります。')
heading(d,'11. データアクセス・トランザクション・セキュリティ'); table(d,['項目','要件'],[['接続アカウント','通常のプロジェクト用アカウントを使用し、SYS/SYSTEMを対象外とする。'],['DDL実行','Oracle DDLは自動COMMITされるため、途中エラー後に無条件再実行しない。'],['再実行','00_check_environment.sqlで既存JL_オブジェクトを確認する。'],['認証情報','会員パスワードはアプリケーション側でBCryptを利用し、平文を投入しない。'],['トークン','メール認証・パスワード再設定を専用テーブルで管理する。'],['監査','会員監査と通報処理監査を業務データから分離する。']])
heading(d,'12. メールトークン・画像・ファイル制約'); table(d,['対象','保持内容','確認SQL'],[['メール認証','会員、トークン、期限、使用状態','10_add_email_verification.sql'],['パスワード再設定','会員、トークン、期限、使用状態','11_add_password_reset.sql'],['コミュニティ画像','投稿との関連、ファイル情報','14_add_community_image.sql、25_create_community_user_features.sql'],['本文・理由','投稿・コメント・通報理由等の文字列長制約','01_create_tables.sql'],['ビュー数','コミュニティ投稿の閲覧数','16_add_community_view_count.sql']]); para(d,'画像保存方式、メール送信サーバー、アップロード容量の運用値はSQLだけでは確定できません。')
heading(d,'13. ERD（論理構成）'); table(d,['親エンティティ','子エンティティ','関係'],[['REGION','CATEGORY / PROCEDURE / OFFICE','地域を基準に分類'],['PROCEDURE','DOCUMENT / STEP / SOURCE / PROC_OFFICE','手続に複数の案内要素'],['MEMBER','USER_TASK / POST / NOTIFICATION / AUDIT','会員を基準に利用履歴'],['USER_TASK','USER_CHECK / USER_TASK_EVENT','保存手続の進捗'],['COMMUNITY_POST','COMMENT / LIKE / IMAGE / REPORT / BOOKMARK','投稿中心の機能'],['REPORT','REPORT_AUDIT','通報状態と運営履歴']])
heading(d,'14. SQLファイル一覧（全ファイル）'); table(d,['No.','ファイル','役割'],[[i,f,purpose.get(f,'内容確認が必要')] for i,f in enumerate(files,1)])
heading(d,'15. 初回インストールと更新手順'); table(d,['順序','実行対象','目的'],[['0','00_check_environment.sql','環境確認'],['1','01_create_tables.sql','基本DDL作成'],['2','02_sample_data.sql','基本データ投入'],['3','04_seed_translation_data.sql','翻訳投入'],['4','07_update_document_guidance.sql','書類案内投入・更新'],['5','08_add_resident_procedures.sql','生活手続追加'],['6','09_review_resident_procedure_documents.sql','書類・翻訳の補完'],['7','03_verify.sql','生成結果確認'],['追加','10〜25の個別SQL','既存環境を確認して必要なものだけ適用']]); para(d,'install_all.sqlは初回インストール順をまとめた実行ファイルです。追加・修正SQLを同じアカウントで全て無条件に実行する運用ではありません。')
heading(d,'16. 検証項目'); table(d,['確認対象','確認内容','根拠'],[['オブジェクト','JL_テーブル、シーケンス、インデックスの存在','00_check_environment.sql'],['件数','地域、カテゴリ、手続、書類、手順等の投入','03_verify.sql'],['整合性','外部キーの親データ存在','01と各投入SQL'],['状態値','CHECK制約にない値がないか','01_create_tables.sql'],['多言語','ja・ko・enの必要案内','04、18〜24'],['通報監査','通報状態と履歴の分離','25_create_community_user_features.sql']])
heading(d,'17. 残存課題と注意事項'); table(d,['優先度','課題','対応方針'],[['高','Oracleでの実行結果は本書作成時点では未確認。','実行前にSQL全文、バックアップ、環境確認を行う。'],['高','初回DDLと追加・修正SQLが混在。','install_all対象と個別更新対象を分ける。'],['中','Oracle識別子長制限の影響。','SQLに定義された短縮名を正本として使用する。'],['中','画像・メールの実運用設定はSQL外。','アプリケーション設定・運用設計で別途定義する。'],['低','サンプルSQLと本番SQLが同じフォルダ。','本番投入前にサンプル投入SQLを分離確認する。']])
heading(d,'18. 参照資料'); table(d,['資料','確認内容'],[['sql/install_all.sql','初回インストール実行順'],['sql/01_create_tables.sql','現行テーブル・制約・インデックス・シーケンス'],['sql/00_check_environment.sql / 03_verify.sql','環境・生成結果確認'],['sql/02〜25の各SQL','初期データ、翻訳、認証、タスク、コミュニティ、修正履歴'],['docs/database.md','データベース概念とアプリ利用'],['README.md / sql/README_SQL.md','システム構成とSQL実行上の注意']])
f=d.sections[0].footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(f,'Japan Life Guide  |  Database Requirements',8,False,(120,110,100))
d.save(OUT); print(OUT)
